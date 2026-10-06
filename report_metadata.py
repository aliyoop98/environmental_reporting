from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Dict, Iterable, Mapping, Optional

import pandas as pd

TEMPLATE_COLUMNS = [
    "Serial Number",
    "Chart Title",
    "Probe ID",
    "Equipment ID",
    "Materials",
    "Profile",
    "Temperature Lower",
    "Temperature Upper",
    "Humidity Lower",
    "Humidity Upper",
    "Notes",
]

FIELD_ALIASES: Dict[str, tuple[str, ...]] = {
    "Serial Number": (
        "serial number", "serial", "serial no", "serial #", "serial id", "device id", "device serial"
    ),
    "Chart Title": (
        "chart title", "title", "report title", "room", "room name", "location", "location name", "space name"
    ),
    "Probe ID": (
        "probe id", "probe", "probe number", "probe #", "probe label"
    ),
    "Equipment ID": (
        "equipment id", "equipment", "equipment number", "equipment #", "equipment name"
    ),
    "Materials": (
        "materials", "material", "materials contents", "materials / contents", "contents", "materials list"
    ),
    "Profile": (
        "profile", "environmental profile", "environment profile", "storage profile"
    ),
    "Temperature Lower": (
        "temperature lower", "temperature min", "temp lower", "temp min", "temperature low", "temp low"
    ),
    "Temperature Upper": (
        "temperature upper", "temperature max", "temp upper", "temp max", "temperature high", "temp high"
    ),
    "Humidity Lower": (
        "humidity lower", "humidity min", "rh lower", "rh min", "humidity low", "rh low"
    ),
    "Humidity Upper": (
        "humidity upper", "humidity max", "rh upper", "rh max", "humidity high", "rh high"
    ),
    "Notes": ("notes", "note", "comments", "comment"),
}


def _header_key(value: object) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("#", " number ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def _serial_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text


def _clean_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null"} else text


def _find_alias(columns: Iterable[object], canonical: str) -> Optional[object]:
    alias_keys = {_header_key(alias) for alias in FIELD_ALIASES[canonical]}
    for col in columns:
        if _header_key(col) in alias_keys:
            return col
    return None


def normalize_metadata_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Return one normalized metadata row per serial number.

    Header aliases are accepted so existing user spreadsheets can often be used
    directly. Blank optional cells mean "use app/automatic default".
    """

    if df is None or df.empty:
        return pd.DataFrame(columns=TEMPLATE_COLUMNS)

    source = df.copy()
    rename_map: Dict[object, str] = {}
    used = set()
    for canonical in TEMPLATE_COLUMNS:
        match = _find_alias([c for c in source.columns if c not in used], canonical)
        if match is not None:
            rename_map[match] = canonical
            used.add(match)

    source = source.rename(columns=rename_map)
    if "Serial Number" not in source.columns:
        raise ValueError(
            "Metadata file needs a Serial Number column (aliases such as Serial, Serial #, or Device ID are accepted)."
        )

    for column in TEMPLATE_COLUMNS:
        if column not in source.columns:
            source[column] = ""

    source = source[TEMPLATE_COLUMNS].copy()
    source["Serial Number"] = source["Serial Number"].map(_serial_text)
    source = source[source["Serial Number"] != ""].copy()

    text_columns = ["Chart Title", "Probe ID", "Equipment ID", "Materials", "Profile", "Notes"]
    for column in text_columns:
        source[column] = source[column].map(_clean_text)

    numeric_columns = [
        "Temperature Lower",
        "Temperature Upper",
        "Humidity Lower",
        "Humidity Upper",
    ]
    for column in numeric_columns:
        source[column] = pd.to_numeric(source[column], errors="coerce")

    # Last row wins for duplicate serials. This makes it easy to append a revised
    # mapping at the bottom of a maintained sheet without deleting history first.
    source = source.drop_duplicates(subset=["Serial Number"], keep="last")
    return source.reset_index(drop=True)


def read_metadata_bytes(file_name: str, data: bytes) -> pd.DataFrame:
    suffix = Path(file_name or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str, engine="openpyxl")
        if not sheets:
            return pd.DataFrame(columns=TEMPLATE_COLUMNS)
        if "Metadata" in sheets:
            raw = sheets["Metadata"]
        else:
            raw = next(iter(sheets.values()))
    elif suffix == ".csv":
        raw = pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False)
    else:
        raise ValueError("Metadata mapping must be a .csv, .xlsx, or .xlsm file.")
    return normalize_metadata_dataframe(raw)


def metadata_lookup(df: pd.DataFrame) -> Dict[str, Dict[str, object]]:
    if df is None or df.empty:
        return {}
    result: Dict[str, Dict[str, object]] = {}
    for _, row in df.iterrows():
        serial = _serial_text(row.get("Serial Number"))
        if not serial:
            continue
        result[serial] = {column: row.get(column) for column in TEMPLATE_COLUMNS if column != "Serial Number"}
    return result


def template_dataframe(serial_data: Optional[Mapping[str, Mapping[str, object]]] = None) -> pd.DataFrame:
    rows = []
    if serial_data:
        for serial in sorted(serial_data, key=lambda x: str(x).lower()):
            info = serial_data[serial]
            rows.append(
                {
                    "Serial Number": str(serial),
                    "Chart Title": str(info.get("default_label") or info.get("option_label") or ""),
                    "Probe ID": "",
                    "Equipment ID": "",
                    "Materials": "",
                    "Profile": "",
                    "Temperature Lower": "",
                    "Temperature Upper": "",
                    "Humidity Lower": "",
                    "Humidity Upper": "",
                    "Notes": "",
                }
            )
    else:
        rows = [
            {
                "Serial Number": "250000001",
                "Chart Title": "Olympus Scanner Room 110",
                "Probe ID": "234",
                "Equipment ID": "Olympus Scanner 228",
                "Materials": "Processed patient slides",
                "Profile": "Olympus",
                "Temperature Lower": "15",
                "Temperature Upper": "28",
                "Humidity Lower": "0",
                "Humidity Upper": "80",
                "Notes": "Example row - replace with your serial mapping",
            }
        ]
    return pd.DataFrame(rows, columns=TEMPLATE_COLUMNS)


def limits_from_metadata(row: Mapping[str, object]) -> Dict[str, tuple[float, float]]:
    result: Dict[str, tuple[float, float]] = {}
    for channel, low_key, high_key in (
        ("Temperature", "Temperature Lower", "Temperature Upper"),
        ("Humidity", "Humidity Lower", "Humidity Upper"),
    ):
        low = pd.to_numeric(pd.Series([row.get(low_key)]), errors="coerce").iloc[0]
        high = pd.to_numeric(pd.Series([row.get(high_key)]), errors="coerce").iloc[0]
        if pd.notna(low) and pd.notna(high) and float(low) < float(high):
            result[channel] = (float(low), float(high))
    return result


def template_workbook_bytes(serial_data: Optional[Mapping[str, Mapping[str, object]]] = None) -> bytes:
    """Build a user-friendly XLSX metadata template."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = template_dataframe(serial_data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    input_font = Font(color="0000FF")
    note_fill = PatternFill("solid", fgColor="FFF2CC")

    for col_idx, column in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(data.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value if value != "" else None)
            cell.font = input_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TEMPLATE_COLUMNS))}{max(2, len(data) + 1)}"
    widths = {
        "A": 18, "B": 34, "C": 14, "D": 28, "E": 34, "F": 20,
        "G": 20, "H": 20, "I": 18, "J": 18, "K": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 30

    instructions = wb.create_sheet("Instructions")
    instructions["A1"] = "Environmental Reporting metadata template"
    instructions["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    instructions["A1"].fill = header_fill
    instructions.merge_cells("A1:D1")
    instructions["A3"] = "Required"
    instructions["B3"] = "Serial Number"
    instructions["A4"] = "Defaults"
    instructions["B4"] = "Chart Title, Probe ID, Equipment ID, Materials"
    instructions["A5"] = "Optional"
    instructions["B5"] = "Profile and Temperature/Humidity limits. Leave blank to use app auto-detection."
    instructions["A6"] = "Customization"
    instructions["B6"] = "Uploaded values become defaults only. You can still edit any report detail or limit in Streamlit."
    instructions["A7"] = "Duplicate serials"
    instructions["B7"] = "If a serial appears more than once, the last row is used."
    instructions["A9"] = "Supported profile names"
    instructions["B9"] = "Room, Olympus, Olympus Room, Fridge, Freezer, Freezer -80 (or leave blank/Auto)."
    instructions["A11"] = "Header aliases"
    instructions["B11"] = "Existing sheets can use common names such as Serial, Device ID, Title, Probe, Equipment, Material, Temp Min/Max, and RH Min/Max."
    for row in range(3, 12):
        instructions[f"A{row}"].font = Font(bold=True, color="666666")
        instructions[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    instructions["A13"] = "Tip"
    instructions["A13"].fill = note_fill
    instructions["A13"].font = Font(bold=True)
    instructions["B13"] = "For the fastest setup, upload your Traceable files first, then download the detected-serial template from the app. It will already contain every detected serial number."
    instructions["B13"].fill = note_fill
    instructions["B13"].alignment = Alignment(wrap_text=True)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 95
    instructions.sheet_view.showGridLines = False
    ws.sheet_view.showGridLines = False

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
